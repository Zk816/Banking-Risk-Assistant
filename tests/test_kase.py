from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from fin_assistant.ingest.kase import _BIN, KaseError, parse_700n, to_metrics

HEADER = [
    "RNUM",
    "REPORT_DATE",
    "CREDITOR_ID",
    "Номер счета",
    "Наименование номера счета",
    "Признак резидентства",
    "Код сектора экономики",
    "Код группы валют",
    "Сумма",
]
D = datetime(2026, 6, 30)


def _sheet(tmp_path: Path, rows: list[tuple[str, str, float]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    for i, (acc, name, amount) in enumerate(rows, 1):
        ws.append([i, D, 14, acc, name, "1", "3", "3", amount])
    path = tmp_path / "x_700h.xlsx"
    wb.save(path)
    return path


ROWS = [
    ("1001", "Наличность в кассе", 100.0),
    ("1411", "Краткосрочные займы клиентам", 800.0),
    ("1424", "Просроченная задолженность клиентов по займам", 50.0),
    ("1428", "Резервы (провизии) по займам", -40.0),
    ("2206", "Краткосрочные вклады физических лиц", 600.0),
    ("2227", "Обязательства по аренде", 10.0),
    ("3599", "Нераспределенная чистая прибыль", 30.0),
    ("3025", "Уставный капитал", 270.0),
]


def test_parses_balances(tmp_path: Path) -> None:
    balances = parse_700n(_sheet(tmp_path, ROWS))
    assert len(balances) == len(ROWS)
    assert balances[0].report_date == D.date()
    assert balances[0].account == "1001"


def test_rows_of_one_account_are_summed(tmp_path: Path) -> None:
    rows = ROWS + [("1001", "Наличность в кассе", 5.0)]
    m = to_metrics(parse_700n(_sheet(tmp_path, rows)))
    assert m["assets"] == 100 + 5 + 800 + 50 - 40


def test_metric_rules(tmp_path: Path) -> None:
    m = to_metrics(parse_700n(_sheet(tmp_path, ROWS)))
    assert m["loans_gross"] == 850
    assert m["loans_net"] == 810
    assert m["overdue_loans"] == 50
    assert m["credit_loss_allowance"] == 40
    assert m["deposits"] == 600
    assert m["net_income_ytd"] == 30
    assert m["assets"] == m["liabilities"] + m["equity"]


def test_unknown_layout_is_refused(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.append(["какая-то", "другая", "таблица"])
    path = tmp_path / "old.xlsx"
    wb.save(path)
    with pytest.raises(KaseError):
        parse_700n(path)


def test_legacy_xls_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "old.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(KaseError):
        parse_700n(path)


def test_bin_on_issuer_page() -> None:
    assert _BIN.search("<p>БИН – 940140000385</p>").group(1) == "940140000385"
    assert _BIN.search("БИН - 123456789012").group(1) == "123456789012"


def test_date_from_file_name() -> None:
    from datetime import date

    from fin_assistant.ingest.kase import _date_from_filename

    assert _date_from_filename(Path("eubn_700h_311221.xlsx")) == date(2021, 12, 31)
    assert _date_from_filename(Path("hsbk_700h_010726.xlsx")) == date(2026, 6, 30)
    assert _date_from_filename(Path("asbn_700h_3_2015_rus.xlsx")) is None


def test_regulation_date_in_header_is_not_the_report_date() -> None:
    from datetime import date

    from fin_assistant.ingest.kase import _report_date

    rows = [
        ("", "Приложение 2 к постановлению Правления от 21.04.2020г. №54"),
        ("", "Отчет об остатках"),
        ("за 30 сентября 2024 года", ""),
    ]
    assert _report_date(rows) == date(2024, 9, 30)


def test_legacy_codes_and_slash_dates() -> None:
    from datetime import date

    from fin_assistant.ingest.kase import _cell, _report_date

    assert _cell(1000.0) == "1000"
    assert _cell(1.0) == "1"
    assert _report_date([("",), (" на 31/ 03/ 2006 года ",)]) == date(2006, 3, 31)


def test_unit_guess_for_report_without_unit_marker(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [
        ["Сведения по остаткам"],
        ["на 31/ 03/ 2006 года"],
        ["КОДЫ", "Наименование", "Сумма"],
        [1, "АКТИВЫ", 1000],
        [1001, "Наличность", 1000],
        [2, "ОБЯЗАТЕЛЬСТВА", 800],
        [2203, "Текущие счета", 800],
        [3, "КАПИТАЛ", 200],
        [3599, "Прибыль", 200],
    ]:
        ws.append(row)
    path = tmp_path / "old_700h.xlsx"
    wb.save(path)
    m = to_metrics(parse_700n(path))
    assert m["assets"] == 1_000_000
    assert m["assets"] == m["liabilities"] + m["equity"]
